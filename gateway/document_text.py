"""Helpers for extracting readable text from local and Juhe remote documents."""

from __future__ import annotations

from dataclasses import dataclass
import asyncio
import logging
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any, Optional
from urllib.parse import urlsplit, urlunsplit
import xml.etree.ElementTree as ET
import zipfile

logger = logging.getLogger(__name__)

_WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
_SPREADSHEET_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".log", ".json", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg"}
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PDF_MIME = "application/pdf"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLS_MIME = "application/vnd.ms-excel"
_MAX_DOCUMENT_EXTRACT_BYTES = 5 * 1024 * 1024
_MAX_DOCUMENT_EXTRACT_CHARS = 20_000
_MAX_PDF_PAGES = 20
_REMOTE_DOCUMENT_TIMEOUT_SECONDS = 90
_REMOTE_MARKDOWN_FETCH_TIMEOUT_SECONDS = 30
_MINERU_LIGHTWEIGHT_ENDPOINT = "https://mineru.net/api/v1/agent/parse/url"
_MINERU_LIGHTWEIGHT_QUERY_ENDPOINT = "https://mineru.net/api/v1/agent/parse"
_MINERU_PRECISE_ENDPOINT = "https://mineru.net/api/v4/extract/task/batch"
_MINERU_PRECISE_QUERY_ENDPOINT = "https://mineru.net/api/v4/extract/task"
_MINERU_POLL_INTERVAL_SECONDS = 2.0
_MINERU_POLL_TIMEOUT_SECONDS = 45.0
_MINERU_CACHE_TOLERANCE_SECONDS = 900
_MULTISPACE_RE = re.compile(r"[ \t]{2,}")
_BLANK_LINE_RE = re.compile(r"\n{3,}")
_REMOTE_DOCUMENT_MIME_TYPES = {
    _PDF_MIME,
    _DOCX_MIME,
    "application/msword",
    _XLSX_MIME,
    _XLS_MIME,
}


@dataclass(frozen=True)
class ExtractedDocumentText:
    display_name: str
    text: str
    truncated: bool = False


@dataclass(frozen=True)
class JuheDocumentParseBundle:
    display_name: str
    extracted_text: str
    display_description: str
    parser: str
    content_format: str
    parse_status: str = "success"
    error_message: str = ""


@dataclass(frozen=True)
class RemoteDocumentAgentContext:
    display_name: str
    extracted_text: str
    display_description: str
    agent_context: str
    parser: str
    from_cache: bool = False


def build_document_context_for_agent(cached_path: str, media_type: str) -> str:
    extracted = extract_local_document_text(cached_path, media_type)
    if extracted is None:
        return ""

    note = (
        f"[The user sent a document: '{extracted.display_name}'. "
        f"Its extracted content is included below. "
        f"The file is also saved at: {cached_path}]"
    )
    return f"{note}\n\n[Content of {extracted.display_name}]:\n{extracted.text}"


def _truncate_display_description(text: str, limit: int = 240) -> str:
    normalized = _normalize_extracted_text(text).replace("\n", " ").strip()
    if len(normalized) <= limit:
        return normalized
    return normalized[: max(limit - 1, 1)].rstrip() + "…"


def _build_document_display_description(display_name: str, extracted_text: str) -> str:
    normalized = _normalize_extracted_text(extracted_text)
    if not normalized:
        return f"文件 {display_name}：已收到，但暂未提取到正文"

    parts = [line.strip() for line in normalized.splitlines() if line.strip()]
    summary = "；".join(parts[:4]) if parts else normalized
    summary = _truncate_display_description(summary)
    return f"文件 {display_name}：{summary}"


def _build_remote_document_context_bundle(
    *,
    display_name: str,
    text: str,
    parser: str,
    from_cache: bool,
    display_description: str | None = None,
) -> RemoteDocumentAgentContext:
    resolved_description = display_description or _build_document_display_description(display_name, text)
    agent_context = (
        f"{resolved_description}\n\n"
        f"[The user sent a document: '{display_name}'. "
        f"Remote-extracted text is included below.]\n\n"
        f"[Content of {display_name}]:\n{text}"
    )
    return RemoteDocumentAgentContext(
        display_name=display_name,
        extracted_text=text,
        display_description=resolved_description,
        agent_context=agent_context,
        parser=parser,
        from_cache=from_cache,
    )


def _build_remote_document_failure_context(
    display_name: str,
    *,
    display_description: str | None = None,
    parser: str = "remote_document",
) -> RemoteDocumentAgentContext:
    resolved_description = display_description or f"文件 {display_name}：已收到，但暂未提取到正文"
    agent_context = (
        f"{resolved_description}\n\n"
        f"[The user sent a document: '{display_name}'. "
        f"Hermes couldn't extract readable text from the remote document URL, "
        f"so only the file metadata is available.]"
    )
    return RemoteDocumentAgentContext(
        display_name=display_name,
        extracted_text="",
        display_description=resolved_description,
        agent_context=agent_context,
        parser=parser,
        from_cache=False,
    )


def _build_parse_bundle(
    *,
    display_name: str,
    extracted_text: str,
    parser: str,
    content_format: str = "text",
) -> JuheDocumentParseBundle:
    return JuheDocumentParseBundle(
        display_name=display_name,
        extracted_text=extracted_text,
        display_description=_build_document_display_description(display_name, extracted_text),
        parser=parser,
        content_format=content_format,
        parse_status="success",
    )


def _build_failed_bundle(
    *,
    display_name: str,
    parser: str,
    error_message: str = "",
    display_description: str | None = None,
) -> JuheDocumentParseBundle:
    return JuheDocumentParseBundle(
        display_name=display_name,
        extracted_text="",
        display_description=display_description or f"文件 {display_name}：已收到，但暂未提取到正文",
        parser=parser,
        content_format="text",
        parse_status="failed",
        error_message=str(error_message or "").strip(),
    )


def _bundle_from_cache_record(record: dict[str, Any], display_name: str) -> JuheDocumentParseBundle:
    resolved_name = str(display_name or record.get("file_name") or "document").strip() or "document"
    return JuheDocumentParseBundle(
        display_name=resolved_name,
        extracted_text=str(record.get("extracted_text") or ""),
        display_description=str(record.get("display_description") or "").strip()
        or _build_document_display_description(resolved_name, str(record.get("extracted_text") or "")),
        parser=str(record.get("parser") or "cache"),
        content_format=str(record.get("content_format") or "text"),
        parse_status=str(record.get("parse_status") or "success").strip().lower() or "success",
        error_message=str(record.get("error_message") or ""),
    )


def extract_local_document_text(cached_path: str, media_type: str) -> Optional[ExtractedDocumentText]:
    path = Path(cached_path)
    if not path.is_file():
        return None

    try:
        if path.stat().st_size > _MAX_DOCUMENT_EXTRACT_BYTES:
            return None
    except OSError:
        return None

    normalized_media = str(media_type or "").strip().lower()
    ext = path.suffix.lower()

    try:
        if ext in _TEXT_EXTENSIONS or normalized_media.startswith("text/"):
            extracted = path.read_text(encoding="utf-8")
        elif ext == ".docx" or normalized_media == _DOCX_MIME:
            extracted = _extract_docx_text(path)
        elif ext == ".doc":
            extracted = _extract_legacy_word_text(path)
        elif ext == ".pdf" or normalized_media == _PDF_MIME:
            extracted = _extract_pdf_text(path)
        elif ext == ".xlsx" or normalized_media == _XLSX_MIME:
            extracted = _extract_xlsx_text(path)
        elif ext == ".xls" or normalized_media == _XLS_MIME:
            extracted = _extract_xls_text(path)
        else:
            return None
    except Exception:
        logger.warning("Failed to extract document text from %s", cached_path, exc_info=True)
        return None

    normalized, truncated = _truncate_document_text(extracted)
    if not normalized:
        return None

    return ExtractedDocumentText(display_name=path.name, text=normalized, truncated=truncated)


def _truncate_document_text(text: str) -> tuple[str, bool]:
    normalized = _normalize_extracted_text(text)
    if not normalized:
        return "", False
    truncated = False
    if len(normalized) > _MAX_DOCUMENT_EXTRACT_CHARS:
        normalized = normalized[:_MAX_DOCUMENT_EXTRACT_CHARS].rstrip()
        normalized += "\n\n[Truncated after 20000 characters]"
        truncated = True
    return normalized, truncated


def _coerce_remote_content_to_text(content: Any) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        pieces = []
        for item in content:
            if isinstance(item, str):
                pieces.append(item)
                continue
            if isinstance(item, dict):
                text = item.get("text")
                if text:
                    pieces.append(str(text))
        return "\n".join(piece for piece in pieces if piece)
    return str(content or "")


def _display_name_from_reference(reference: str, fallback: str = "document") -> str:
    text = str(reference or "").strip()
    if not text:
        return fallback
    parsed = urlsplit(text)
    if parsed.scheme and parsed.netloc:
        derived = Path(parsed.path).name
        return derived or fallback
    return Path(text).name or fallback


def _supports_remote_document_url(media_type: str, display_name: str) -> bool:
    normalized_media = str(media_type or "").strip().lower()
    if normalized_media in _REMOTE_DOCUMENT_MIME_TYPES:
        return True
    return Path(display_name).suffix.lower() in {".pdf", ".docx", ".doc", ".xlsx", ".xls"}


def _dashscope_document_endpoint_from_base_url(base_url: str) -> str:
    text = str(base_url or "").strip()
    if not text:
        return ""
    try:
        parsed = urlsplit(text)
    except ValueError:
        return ""
    if "dashscope" not in parsed.netloc.lower():
        return ""
    scheme = parsed.scheme or "https"
    return urlunsplit(
        (
            scheme,
            parsed.netloc,
            "/api/v1/services/aigc/text-generation/generation",
            "",
            "",
        )
    )


def _resolve_dashscope_document_runtime() -> Optional[dict[str, str]]:
    from hermes_cli.runtime_provider import resolve_runtime_provider

    requested_provider = os.getenv("HERMES_INFERENCE_PROVIDER")
    try:
        runtime = resolve_runtime_provider(requested=requested_provider)
    except Exception:
        logger.warning("Failed to resolve runtime provider for remote document parsing", exc_info=True)
        return None

    api_key = str(runtime.get("api_key") or "").strip()
    base_url = str(runtime.get("base_url") or "").strip()
    endpoint = _dashscope_document_endpoint_from_base_url(base_url)
    if not api_key or not endpoint:
        return None

    provider = str(runtime.get("provider") or "").strip().lower()
    if provider == "alibaba" or "dashscope" in endpoint:
        return {
            "provider": provider,
            "api_key": api_key,
            "endpoint": endpoint,
        }
    return None


async def _extract_document_text_via_dashscope_doc_url(
    *,
    document_url: str,
    display_name: str,
    media_type: str,
) -> Optional[ExtractedDocumentText]:
    runtime = _resolve_dashscope_document_runtime()
    if runtime is None:
        return None
    if not _supports_remote_document_url(media_type, display_name):
        return None

    import httpx

    payload = {
        "model": "qwen-doc-turbo",
        "input": {
            "messages": [
                {
                    "role": "system",
                    "content": (
                        "You extract document text faithfully. "
                        "Return only the readable text in order. "
                        "Do not summarize, explain, or add JSON."
                    ),
                },
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": (
                                "Extract the readable text from this document. "
                                "Preserve the original reading order. "
                                "If the document contains scanned pages, OCR them. "
                                "Return only the extracted text."
                            ),
                        },
                        {
                            "type": "doc_url",
                            "doc_url": [document_url],
                            "file_parsing_strategy": "text_and_images",
                        },
                    ],
                },
            ]
        },
    }
    headers = {
        "Authorization": f"Bearer {runtime['api_key']}",
        "Content-Type": "application/json",
    }

    try:
        async with httpx.AsyncClient(timeout=_REMOTE_DOCUMENT_TIMEOUT_SECONDS) as client:
            response = await client.post(runtime["endpoint"], headers=headers, json=payload)
            response.raise_for_status()
            body = response.json()
    except Exception:
        logger.warning("DashScope remote document parsing failed for %s", display_name, exc_info=True)
        return None

    content = (
        body.get("output", {})
        .get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
    )
    extracted_text = _coerce_remote_content_to_text(content)
    normalized, truncated = _truncate_document_text(extracted_text)
    if not normalized:
        return None
    return ExtractedDocumentText(
        display_name=display_name or _display_name_from_reference(document_url),
        text=normalized,
        truncated=truncated,
    )


async def _download_document_to_temp_path(document_url: str, display_name: str) -> Optional[Path]:
    text = str(document_url or "").strip()
    if not text:
        return None

    local_path = Path(text)
    if local_path.is_file():
        return local_path

    import httpx

    suffix = Path(display_name or _display_name_from_reference(text)).suffix or ".bin"
    fd, tmp_name = tempfile.mkstemp(prefix="hermes-juhe-doc-", suffix=suffix)
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        async with httpx.AsyncClient(timeout=_REMOTE_DOCUMENT_TIMEOUT_SECONDS, follow_redirects=True) as client:
            async with client.stream("GET", text) as response:
                response.raise_for_status()
                with tmp_path.open("wb") as handle:
                    async for chunk in response.aiter_bytes():
                        if chunk:
                            handle.write(chunk)
        return tmp_path
    except Exception:
        tmp_path.unlink(missing_ok=True)
        logger.warning("Failed to download Juhe document %s for local parsing", display_name, exc_info=True)
        return None


async def _fetch_remote_text(url: str) -> str:
    import httpx

    async with httpx.AsyncClient(timeout=_REMOTE_MARKDOWN_FETCH_TIMEOUT_SECONDS, follow_redirects=True) as client:
        response = await client.get(url)
        response.raise_for_status()
        return response.text


def _extract_markdown_url_from_payload(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("markdown_url", "md_url", "full_markdown_url", "result_url", "download_url"):
            value = payload.get(key)
            text = str(value or "").strip()
            if text:
                return text
        for nested_key in ("data", "result", "output"):
            value = payload.get(nested_key)
            nested = _extract_markdown_url_from_payload(value)
            if nested:
                return nested
        files = payload.get("files")
        if isinstance(files, list):
            for item in files:
                nested = _extract_markdown_url_from_payload(item)
                if nested:
                    return nested
    elif isinstance(payload, list):
        for item in payload:
            nested = _extract_markdown_url_from_payload(item)
            if nested:
                return nested
    return ""


def _extract_task_id_from_payload(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("task_id", "taskId", "id"):
            value = str(payload.get(key) or "").strip()
            if value:
                return value
        for nested_key in ("data", "result", "output"):
            value = payload.get(nested_key)
            nested = _extract_task_id_from_payload(value)
            if nested:
                return nested
    elif isinstance(payload, list):
        for item in payload:
            nested = _extract_task_id_from_payload(item)
            if nested:
                return nested
    return ""


def _extract_inline_text_from_payload(payload: Any) -> str:
    if isinstance(payload, dict):
        for key in ("markdown", "md", "text", "content"):
            value = payload.get(key)
            text = _coerce_remote_content_to_text(value)
            if text.strip():
                return text
        for nested_key in ("data", "result", "output"):
            value = payload.get(nested_key)
            nested = _extract_inline_text_from_payload(value)
            if nested:
                return nested
    elif isinstance(payload, list):
        for item in payload:
            nested = _extract_inline_text_from_payload(item)
            if nested:
                return nested
    return ""


async def _poll_mineru_markdown_url(
    *,
    query_endpoint: str,
    task_id: str,
    headers: Optional[dict[str, str]] = None,
) -> str:
    import httpx

    deadline = asyncio.get_running_loop().time() + _MINERU_POLL_TIMEOUT_SECONDS
    async with httpx.AsyncClient(timeout=_REMOTE_DOCUMENT_TIMEOUT_SECONDS, follow_redirects=True) as client:
        while asyncio.get_running_loop().time() < deadline:
            response = await client.get(f"{query_endpoint.rstrip('/')}/{task_id}", headers=headers)
            response.raise_for_status()
            body = response.json()
            markdown_url = _extract_markdown_url_from_payload(body)
            if markdown_url:
                return markdown_url
            state_text = jsonish_state = _coerce_remote_content_to_text(
                body.get("state") if isinstance(body, dict) else ""
            ).lower()
            if not state_text and isinstance(body, dict):
                state_text = _coerce_remote_content_to_text(body.get("status")).lower()
            if state_text in {"done", "success", "completed", "failed", "error"}:
                break
            await asyncio.sleep(_MINERU_POLL_INTERVAL_SECONDS)
    return ""


async def _extract_document_text_via_mineru_lightweight_url(
    *,
    document_url: str,
    display_name: str,
    media_type: str,
) -> Optional[ExtractedDocumentText]:
    if not _supports_remote_document_url(media_type, display_name):
        return None

    import httpx

    payload = {
        "url": document_url,
        "no_cache": False,
        "cache_tolerance": _MINERU_CACHE_TOLERANCE_SECONDS,
    }
    try:
        async with httpx.AsyncClient(timeout=_REMOTE_DOCUMENT_TIMEOUT_SECONDS, follow_redirects=True) as client:
            response = await client.post(_MINERU_LIGHTWEIGHT_ENDPOINT, json=payload)
            response.raise_for_status()
            body = response.json()
    except Exception:
        logger.warning("MinerU lightweight parse failed for %s", display_name, exc_info=True)
        return None

    markdown_url = _extract_markdown_url_from_payload(body)
    if not markdown_url:
        task_id = _extract_task_id_from_payload(body)
        if task_id:
            markdown_url = await _poll_mineru_markdown_url(
                query_endpoint=_MINERU_LIGHTWEIGHT_QUERY_ENDPOINT,
                task_id=task_id,
            )

    extracted_text = ""
    if markdown_url:
        try:
            extracted_text = await _fetch_remote_text(markdown_url)
        except Exception:
            logger.warning("MinerU lightweight markdown fetch failed for %s", display_name, exc_info=True)
            extracted_text = ""
    if not extracted_text:
        extracted_text = _extract_inline_text_from_payload(body)

    normalized, truncated = _truncate_document_text(extracted_text)
    if not normalized:
        return None
    return ExtractedDocumentText(
        display_name=display_name or _display_name_from_reference(document_url),
        text=normalized,
        truncated=truncated,
    )


async def _extract_document_text_via_mineru_precise_url(
    *,
    document_url: str,
    display_name: str,
    media_type: str,
) -> Optional[ExtractedDocumentText]:
    token = str(os.getenv("MINERU_API_TOKEN", "")).strip()
    if not token or not _supports_remote_document_url(media_type, display_name):
        return None

    import httpx

    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
        "Accept": "*/*",
    }
    payload = {
        "files": [
            {
                "url": document_url,
                "data_id": Path(display_name or "document").stem[:64] or "juhe-doc",
            }
        ],
        "model_version": "vlm",
        "no_cache": False,
        "cache_tolerance": _MINERU_CACHE_TOLERANCE_SECONDS,
    }
    try:
        async with httpx.AsyncClient(timeout=_REMOTE_DOCUMENT_TIMEOUT_SECONDS, follow_redirects=True) as client:
            response = await client.post(_MINERU_PRECISE_ENDPOINT, headers=headers, json=payload)
            response.raise_for_status()
            body = response.json()
    except Exception:
        logger.warning("MinerU precise parse failed for %s", display_name, exc_info=True)
        return None

    markdown_url = _extract_markdown_url_from_payload(body)
    if not markdown_url:
        task_id = _extract_task_id_from_payload(body)
        if task_id:
            markdown_url = await _poll_mineru_markdown_url(
                query_endpoint=_MINERU_PRECISE_QUERY_ENDPOINT,
                task_id=task_id,
                headers=headers,
            )

    extracted_text = ""
    if markdown_url:
        try:
            extracted_text = await _fetch_remote_text(markdown_url)
        except Exception:
            logger.warning("MinerU precise markdown fetch failed for %s", display_name, exc_info=True)
            extracted_text = ""
    if not extracted_text:
        extracted_text = _extract_inline_text_from_payload(body)

    normalized, truncated = _truncate_document_text(extracted_text)
    if not normalized:
        return None
    return ExtractedDocumentText(
        display_name=display_name or _display_name_from_reference(document_url),
        text=normalized,
        truncated=truncated,
    )


async def parse_juhe_document_bundle(
    *,
    document_url: str,
    display_name: str,
    media_type: str,
) -> JuheDocumentParseBundle:
    resolved_name = display_name or _display_name_from_reference(document_url, fallback="document")
    suffix = Path(resolved_name).suffix.lower()
    temp_path: Optional[Path] = None

    try:
        temp_path = await _download_document_to_temp_path(document_url, resolved_name)
        if temp_path is not None:
            local = extract_local_document_text(str(temp_path), media_type)
            if local is not None and local.text.strip():
                content_format = "markdown" if suffix in {".xlsx", ".xls"} else "text"
                return _build_parse_bundle(
                    display_name=resolved_name,
                    extracted_text=local.text,
                    parser=f"local_{suffix.lstrip('.') or 'document'}",
                    content_format=content_format,
                )

        remote: Optional[ExtractedDocumentText]
        if suffix == ".doc" or str(media_type or "").strip().lower() == "application/msword":
            remote = await _extract_document_text_via_mineru_precise_url(
                document_url=document_url,
                display_name=resolved_name,
                media_type=media_type,
            )
            if remote is not None:
                return _build_parse_bundle(
                    display_name=remote.display_name,
                    extracted_text=remote.text,
                    parser="mineru_precise_url",
                    content_format="markdown",
                )
        else:
            remote = await _extract_document_text_via_mineru_lightweight_url(
                document_url=document_url,
                display_name=resolved_name,
                media_type=media_type,
            )
            if remote is not None:
                return _build_parse_bundle(
                    display_name=remote.display_name,
                    extracted_text=remote.text,
                    parser="mineru_lightweight_url",
                    content_format="markdown",
                )

        dashscope = await _extract_document_text_via_dashscope_doc_url(
            document_url=document_url,
            display_name=resolved_name,
            media_type=media_type,
        )
        if dashscope is not None:
            return _build_parse_bundle(
                display_name=dashscope.display_name,
                extracted_text=dashscope.text,
                parser="dashscope_doc_url",
                content_format="text",
            )
    except Exception as exc:
        logger.warning("Juhe document parsing failed for %s", resolved_name, exc_info=True)
        return _build_failed_bundle(
            display_name=resolved_name,
            parser="juhe_document_parse",
            error_message=str(exc),
        )
    finally:
        if temp_path is not None and temp_path.is_file() and not str(document_url or "").strip() == str(temp_path):
            temp_path.unlink(missing_ok=True)

    return _build_failed_bundle(
        display_name=resolved_name,
        parser="juhe_document_parse",
    )


async def build_remote_document_context_for_agent(
    *,
    media_ref: str,
    document_url: str,
    media_type: str,
    display_name: str = "",
    session_db: Any = None,
    cache_identity: Optional[dict[str, Any]] = None,
) -> RemoteDocumentAgentContext:
    normalized_name = display_name or _display_name_from_reference(media_ref, fallback="document")

    cached = None
    if session_db is not None and cache_identity and hasattr(session_db, "get_juhe_attachment_parse_cache"):
        cached = session_db.get_juhe_attachment_parse_cache(platform="juhe", **cache_identity)
    cached_bundle = _bundle_from_cache_record(cached, normalized_name) if cached else None
    if cached_bundle and cached_bundle.parse_status == "success" and cached_bundle.extracted_text.strip():
        return _build_remote_document_context_bundle(
            display_name=normalized_name,
            text=cached_bundle.extracted_text,
            parser=cached_bundle.parser,
            from_cache=True,
            display_description=cached_bundle.display_description,
        )

    if not document_url or not _supports_remote_document_url(media_type, normalized_name):
        return _build_remote_document_failure_context(
            normalized_name,
            display_description=cached_bundle.display_description if cached_bundle else None,
            parser=cached_bundle.parser if cached_bundle else "juhe_document_parse",
        )

    bundle = await parse_juhe_document_bundle(
        document_url=document_url,
        display_name=normalized_name,
        media_type=media_type,
    )

    if session_db is not None and cache_identity and hasattr(session_db, "upsert_juhe_attachment_parse_cache"):
        try:
            session_db.upsert_juhe_attachment_parse_cache(
                platform="juhe",
                media_type=media_type,
                file_name=bundle.display_name,
                parser=bundle.parser,
                extracted_text=bundle.extracted_text,
                content_kind="document",
                content_format=bundle.content_format,
                display_description=bundle.display_description,
                parse_status=bundle.parse_status,
                error_message=bundle.error_message,
                **cache_identity,
            )
        except Exception:
            logger.warning("Failed to persist Juhe document parse cache for %s", normalized_name, exc_info=True)

    if bundle.parse_status == "success" and bundle.extracted_text.strip():
        return _build_remote_document_context_bundle(
            display_name=bundle.display_name,
            text=bundle.extracted_text,
            parser=bundle.parser,
            from_cache=False,
            display_description=bundle.display_description,
        )

    fallback_description = bundle.display_description
    if cached_bundle and cached_bundle.display_description:
        fallback_description = cached_bundle.display_description
    return _build_remote_document_failure_context(
        normalized_name,
        display_description=fallback_description,
        parser=bundle.parser,
    )


def _extract_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")

    root = ET.fromstring(document_xml)
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:body/w:p", _WORD_NS):
        line = "".join(node.text or "" for node in paragraph.findall(".//w:t", _WORD_NS)).strip()
        if line:
            paragraphs.append(line)

    if paragraphs:
        return "\n".join(paragraphs)

    fallback = [node.text or "" for node in root.findall(".//w:t", _WORD_NS) if (node.text or "").strip()]
    return "\n".join(item.strip() for item in fallback if item.strip())


def _extract_legacy_word_text(path: Path) -> str:
    textutil = shutil.which("textutil")
    if not textutil:
        return ""

    result = subprocess.run(
        [textutil, "-convert", "txt", "-stdout", str(path)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=20,
        check=False,
    )
    if result.returncode != 0:
        return ""
    return result.stdout


def _extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        if index >= _MAX_PDF_PAGES:
            break
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if page_text:
            pages.append(page_text)
        if sum(len(item) for item in pages) >= _MAX_DOCUMENT_EXTRACT_CHARS:
            break
    return "\n\n".join(pages)


def _extract_xlsx_text(path: Path) -> str:
    text = _extract_xlsx_text_via_openpyxl(path)
    if text:
        return text
    return _extract_xlsx_text_from_zip(path)


def _extract_xlsx_text_via_openpyxl(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except Exception:
        return ""

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ""

    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value).strip() for value in row if value not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
    finally:
        close = getattr(workbook, "close", None)
        if callable(close):
            close()
    return "\n".join(line for line in lines if line.strip())


def _extract_xlsx_text_from_zip(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(".//s:si", _SPREADSHEET_NS):
                parts = [node.text or "" for node in item.findall(".//s:t", _SPREADSHEET_NS)]
                shared_strings.append("".join(parts))

        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet_names: list[str] = [
            sheet.attrib.get("name", f"Sheet{index}")
            for index, sheet in enumerate(workbook_root.findall(".//s:sheets/s:sheet", _SPREADSHEET_NS), start=1)
        ]

        lines: list[str] = []
        sheet_paths = sorted(
            name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        for index, sheet_path in enumerate(sheet_paths):
            sheet_title = sheet_names[index] if index < len(sheet_names) else Path(sheet_path).stem
            lines.append(f"[Sheet: {sheet_title}]")
            root = ET.fromstring(archive.read(sheet_path))
            for row in root.findall(".//s:sheetData/s:row", _SPREADSHEET_NS):
                values: list[str] = []
                for cell in row.findall("s:c", _SPREADSHEET_NS):
                    cell_type = cell.attrib.get("t", "")
                    value = cell.findtext("s:v", default="", namespaces=_SPREADSHEET_NS)
                    if cell_type == "s" and value.isdigit():
                        index_value = int(value)
                        if 0 <= index_value < len(shared_strings):
                            values.append(shared_strings[index_value].strip())
                    else:
                        inline = "".join(
                            node.text or ""
                            for node in cell.findall(".//s:is/s:t", _SPREADSHEET_NS)
                        ).strip()
                        if inline:
                            values.append(inline)
                        elif value.strip():
                            values.append(value.strip())
                cleaned = [item for item in values if item]
                if cleaned:
                    lines.append(" | ".join(cleaned))
        return "\n".join(line for line in lines if line.strip())


def _extract_xls_text(path: Path) -> str:
    try:
        import xlrd
    except Exception:
        return ""

    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
    except Exception:
        return ""

    lines: list[str] = []
    try:
        for sheet_name in workbook.sheet_names():
            lines.append(f"[Sheet: {sheet_name}]")
            sheet = workbook.sheet_by_name(sheet_name)
            for row_index in range(sheet.nrows):
                row_values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
                if row_values:
                    lines.append(" | ".join(row_values))
    finally:
        release = getattr(workbook, "release_resources", None)
        if callable(release):
            release()
    return "\n".join(line for line in lines if line.strip())


def _normalize_extracted_text(text: str) -> str:
    normalized_lines = []
    for raw_line in str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = _MULTISPACE_RE.sub(" ", raw_line).strip()
        if line:
            normalized_lines.append(line)
        elif normalized_lines and normalized_lines[-1] != "":
            normalized_lines.append("")

    normalized = "\n".join(normalized_lines).strip()
    return _BLANK_LINE_RE.sub("\n\n", normalized)
